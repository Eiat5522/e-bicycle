"""Generate the Glide / Lamphun Smart Green Mobility pitch deck as .pptx.

Vendor-neutral: no named cloud vendors or frameworks (Supabase, Next.js,
Expo, Google, etc.). Architecture framed generically as a self-hosted /
locally deployable platform to align with government data-residency and
compliance policy.

Content sourced from docs/product-design/product-design-brief.md
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------------------
# Brand palette
# ---------------------------------------------------------------------------
GLIDE_GREEN = RGBColor(0x12, 0xB8, 0x81)
DARK = RGBColor(0x14, 0x2A, 0x27)
SLATE = RGBColor(0x4A, 0x5A, 0x58)
LIGHT = RGBColor(0xF4, 0xF8, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT2 = RGBColor(0x0E, 0x8C, 0x63)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, color, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    _set_fill(sp, color)
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def add_round(slide, x, y, w, h, color, radius=0.08):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = radius
    _set_fill(sp, color)
    sp.shadow.inherit = False
    return sp


def add_ellipse(slide, x, y, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    _set_fill(sp, color)
    sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(4)
        for (text, size, bold, color, *rest) in para:
            italic = rest[0] if rest else False
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = FONT
            r.font.color.rgb = color
    return tb


def bg(slide, color=LIGHT):
    add_rect(slide, 0, 0, SW, SH, color)


def footer(slide, num):
    add_text(slide, Inches(0.5), Inches(7.05), Inches(9), Inches(0.3),
             [[("Lamphun Smart Green Mobility", 9, False, SLATE)]])
    add_text(slide, Inches(11.8), Inches(7.05), Inches(1.0), Inches(0.3),
             [[(str(num), 9, False, SLATE)]], align=PP_ALIGN.RIGHT)


def section_header(slide, kicker, title):
    add_round(slide, Inches(0.5), Inches(0.55), Inches(0.18), Inches(0.5), GLIDE_GREEN)
    add_text(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.4),
             [[(kicker.upper(), 12, True, ACCENT2)]])
    add_text(slide, Inches(0.8), Inches(0.92), Inches(11.8), Inches(0.95),
             [[(title, 31, True, DARK)]])


# ---------------------------------------------------------------------------
# Slide 1 — Title
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
add_ellipse(s, Inches(9.5), Inches(-2.2), Inches(6), Inches(6), ACCENT2)
add_ellipse(s, Inches(11.2), Inches(3.8), Inches(4.5), Inches(4.5), GLIDE_GREEN)
add_ellipse(s, Inches(-1.5), Inches(5.0), Inches(4), Inches(4), RGBColor(0x22, 0x40, 0x3C))

add_text(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.5),
         [[("LAMPHUN SMART GREEN MOBILITY", 16, True, GLIDE_GREEN)]])
add_text(s, Inches(0.9), Inches(2.25), Inches(11.8), Inches(2.0),
         [[("A Public Green Mobility", 46, True, WHITE)],
          [("Operating System", 46, True, WHITE)]], line_spacing=1.0)
add_text(s, Inches(0.95), Inches(4.5), Inches(11), Inches(0.8),
         [[("Keeping a 100-bike public e-bike service running safely,", 17, False, RGBColor(0xC8, 0xE6, 0xDA))],
          [("transparently, and continuously — on infrastructure we control.", 17, False, RGBColor(0xC8, 0xE6, 0xDA))]])
add_text(s, Inches(0.95), Inches(6.5), Inches(11), Inches(0.5),
         [[("Investor & Stakeholder Briefing  ·  MVP target: 31 October 2026", 13, True, GLIDE_GREEN)]])

# ---------------------------------------------------------------------------
# Slide 2 — The Problem / Context
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "The Context", "A public service can't stop when tech fails")
cards = [
    ("Connectivity", "Stations and tourists hit dead zones; apps and payment gateways drop."),
    ("Trust & transparency", "Government and citizens need verifiable service, financial, and ESG records."),
    ("Inclusion", "Residents and foreign tourists must both register, pay, and ride without friction."),
    ("Residency", "Public mobility data must stay on locally controlled, sovereign infrastructure."),
]
cx, cw, gap = Inches(0.8), Inches(2.9), Inches(0.18)
for i, (h, b) in enumerate(cards):
    x = cx + i * (cw + gap)
    card = add_round(s, x, Inches(2.2), cw, Inches(3.7), WHITE, radius=0.06)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_round(s, x + Inches(0.3), Inches(2.55), Inches(0.7), Inches(0.7), GLIDE_GREEN, radius=0.5)
    add_text(s, x + Inches(0.3), Inches(2.62), Inches(0.7), Inches(0.55),
             [[(str(i + 1), 22, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.3), Inches(3.5), cw - Inches(0.6), Inches(0.5),
             [[(h, 18, True, DARK)]])
    add_text(s, x + Inches(0.3), Inches(4.05), cw - Inches(0.6), Inches(1.7),
             [[(b, 13, False, SLATE)]], line_spacing=1.05)
footer(s, 2)

# ---------------------------------------------------------------------------
# Slide 3 — Product Promise
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
add_round(s, Inches(0.5), Inches(0.55), Inches(0.18), Inches(0.5), GLIDE_GREEN)
add_text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.4),
         [[("PRODUCT PROMISE", 12, True, GLIDE_GREEN)]])
add_text(s, Inches(0.8), Inches(0.95), Inches(11.8), Inches(1.9),
         [[("A rider can register, rent, pay, ride, and return", 32, True, WHITE)],
          [("with confidence —", 32, True, WHITE)],
          [("and the service keeps running even when the app, internet,", 24, False, RGBColor(0xC8, 0xE6, 0xDA))],
          [("or payment gateway fails.", 24, False, RGBColor(0xC8, 0xE6, 0xDA))]], line_spacing=1.0)
promise = [
    "Riders: register → rent → pay → ride → return, with bilingual clarity.",
    "Station staff: start rentals, returns, incidents & payment capture offline.",
    "Operators & government: verify exactly what happened, afterward.",
]
for i, p in enumerate(promise):
    y = Inches(3.5) + i * Inches(0.95)
    add_ellipse(s, Inches(0.9), y, Inches(0.34), Inches(0.34), GLIDE_GREEN)
    add_text(s, Inches(1.35), y - Inches(0.02), Inches(10.5), Inches(0.5),
             [[(p, 16, False, RGBColor(0xE6, 0xF2, 0xED))]], anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------------------
# Slide 4 — Design Principles
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "How We Design", "Four principles that keep the service alive")
principles = [
    ("Operation First", "Every critical journey has a primary digital path AND a staff-assisted offline fallback."),
    ("Safety Before Convenience", "Vehicle readiness, battery, inspection & incident handling outrank speed or polish."),
    ("Evidence At Every Handoff", "Rental, return, payment, damage & override flows log IDs, actor, time, photo proof."),
    ("Bilingual & Simple", "Thai + tourist-ready copy; simple MVP that expands into a full system."),
]
col_w = Inches(5.9)
for i, (h, b) in enumerate(principles):
    col, row = i % 2, i // 2
    x = Inches(0.8) + col * (col_w + Inches(0.3))
    y = Inches(2.2) + row * Inches(2.0)
    card = add_round(s, x, y, col_w, Inches(1.75), WHITE, radius=0.05)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_rect(s, x, y, Inches(0.12), Inches(1.75), GLIDE_GREEN)
    add_text(s, x + Inches(0.35), y + Inches(0.25), col_w - Inches(0.6), Inches(0.5),
             [[(h, 17, True, DARK)]])
    add_text(s, x + Inches(0.35), y + Inches(0.8), col_w - Inches(0.6), Inches(0.9),
             [[(b, 13, False, SLATE)]], line_spacing=1.05)
footer(s, 4)

# ---------------------------------------------------------------------------
# Slide 5 — Audiences
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "Who It Serves", "One system, six roles")
roles = [
    ("Tourist rider", "Register fast, understand rules, pay, unlock, return, get help."),
    ("Local resident", "Repeat rental with low friction and clear pricing."),
    ("Station staff", "Start/return rentals, verify payments, inspect, record incidents."),
    ("Technician", "Inspect, maintain, swap batteries, track safety status."),
    ("Operations manager", "Monitor availability, revenue, incidents, reconciliation."),
    ("Government viewer", "Transparent service, financial, ESG & audit reporting — read-only."),
]
col_w = Inches(3.9); row_h = Inches(1.9)
for i, (h, b) in enumerate(roles):
    col, row = i % 3, i // 3
    x = Inches(0.8) + col * (col_w + Inches(0.2))
    y = Inches(2.2) + row * (row_h + Inches(0.15))
    card = add_round(s, x, y, col_w, row_h, WHITE, radius=0.05)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_rect(s, x, y, Inches(0.12), row_h, ACCENT2)
    add_text(s, x + Inches(0.35), y + Inches(0.25), col_w - Inches(0.6), Inches(0.5),
             [[(h, 16, True, DARK)]])
    add_text(s, x + Inches(0.35), y + Inches(0.8), col_w - Inches(0.6), Inches(1.0),
             [[(b, 12.5, False, SLATE)]], line_spacing=1.05)
footer(s, 5)

# ---------------------------------------------------------------------------
# Slide 6 — Product: Rider Experience
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "Product  ·  Rider", "The mobile app riders already use")
feats = [
    ("Auth & profile", "Register with PDPA consent and identity capture."),
    ("Map discovery", "Find nearby bikes, stations, and battery at a glance."),
    ("Unlock & ride", "Scan or tap to release; active-ride recovery built in."),
    ("Wallet & top-up", "Clear prepaid balance with bilingual clarity."),
    ("Eco impact", "See the carbon you avoided on every trip."),
    ("Support", "Help and incident reporting when something goes wrong."),
]
col_w = Inches(3.9); row_h = Inches(1.9)
for i, (h, b) in enumerate(feats):
    col, row = i % 3, i // 3
    x = Inches(0.8) + col * (col_w + Inches(0.2))
    y = Inches(2.2) + row * (row_h + Inches(0.15))
    card = add_round(s, x, y, col_w, row_h, WHITE, radius=0.05)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_ellipse(s, x + Inches(0.3), y + Inches(0.3), Inches(0.55), Inches(0.55), GLIDE_GREEN)
    add_text(s, x + Inches(0.3), y + Inches(0.32), Inches(0.55), Inches(0.5),
             [[("✓", 16, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.35), y + Inches(1.0), col_w - Inches(0.6), Inches(0.5),
             [[(h, 16, True, DARK)]])
    add_text(s, x + Inches(0.35), y + Inches(1.45), col_w - Inches(0.6), Inches(0.4),
             [[(b, 12, False, SLATE)]], line_spacing=1.0)
footer(s, 6)

# ---------------------------------------------------------------------------
# Slide 7 — Product: Offline-First Station Tablet
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
add_round(s, Inches(0.5), Inches(0.55), Inches(0.18), Inches(0.5), GLIDE_GREEN)
add_text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.4),
         [[("PRODUCT  ·  STATION", 12, True, GLIDE_GREEN)]])
add_text(s, Inches(0.8), Inches(0.95), Inches(11.8), Inches(0.9),
         [[("The offline-first staff tablet keeps service moving", 30, True, WHITE)]], line_spacing=1.0)
left = [
    ("Start & return rentals", "Frictionless staff-assisted rental for riders without the app."),
    ("Verify payments", "QR / PromptPay / manual slip reference with reconciliation state."),
    ("Record incidents", "Rental, vehicle, user, status, and photo evidence."),
]
right = [
    ("Manual override", "Permission-sensitive fallback when systems are down."),
    ("Inspection & battery", "Pre-use checks, condition, battery level at return."),
    ("Fallback backup", "Exportable records + spreadsheet backup workflow if cloud is unreachable."),
]
def col_block(x, items):
    for i, (h, b) in enumerate(items):
        y = Inches(2.2) + i * Inches(1.5)
        card = add_round(s, x, y, Inches(5.7), Inches(1.3), RGBColor(0x22, 0x40, 0x3C), radius=0.05)
        add_ellipse(s, x + Inches(0.25), y + Inches(0.32), Inches(0.55), Inches(0.55), GLIDE_GREEN)
        add_text(s, x + Inches(0.25), y + Inches(0.34), Inches(0.55), Inches(0.5),
                 [[("●", 16, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(1.0), y + Inches(0.2), Inches(4.5), Inches(0.5),
                 [[(h, 15, True, WHITE)]])
        add_text(s, x + Inches(1.0), y + Inches(0.65), Inches(4.5), Inches(0.6),
                 [[(b, 12, False, RGBColor(0xC8, 0xE6, 0xDA))]], line_spacing=1.0)
col_block(Inches(0.8), left)
col_block(Inches(6.8), right)

# ---------------------------------------------------------------------------
# Slide 8 — Product: Operations Console
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "Product  ·  Operations", "One console for managers and government")
items = [
    ("Dashboard KPIs", "Fleet availability, utilization, revenue, incidents, maintenance."),
    ("Ride context", "User, wallet, ride replay, and status-event history."),
    ("Reconciliation", "Digital and manual payments tied to Rental ID, reviewed daily."),
    ("Government viewer", "Read-only transparency: service, financial, ESG, audit history."),
    ("Exports", "CSV/report exports for oversight and daily reconciliation."),
    ("Role separation", "Clear operational roles; government viewer cannot mutate data."),
]
col_w = Inches(3.9); row_h = Inches(1.9)
for i, (h, b) in enumerate(items):
    col, row = i % 3, i // 3
    x = Inches(0.8) + col * (col_w + Inches(0.2))
    y = Inches(2.2) + row * (row_h + Inches(0.15))
    card = add_round(s, x, y, col_w, row_h, WHITE, radius=0.05)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_rect(s, x, y, Inches(0.12), row_h, ACCENT2)
    add_text(s, x + Inches(0.35), y + Inches(0.25), col_w - Inches(0.6), Inches(0.5),
             [[(h, 16, True, DARK)]])
    add_text(s, x + Inches(0.35), y + Inches(0.8), col_w - Inches(0.6), Inches(1.0),
             [[(b, 12.5, False, SLATE)]], line_spacing=1.05)
footer(s, 8)

# ---------------------------------------------------------------------------
# Slide 9 — The Fleet (FreeDare)
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "The Fleet", "100 FreeDare e-bikes, built for public service")
specs = [
    ("100 units", "Public fleet across Lamphun at launch."),
    ("48V · 500–750W", "Motor power for hills and payloads."),
    ("Li-ion 20Ah", "Removable battery for swap-and-go ops."),
    ("IPX5 rated", "Weather-resistant for daily outdoor use."),
    ("60–120 km", "Assisted range per charge."),
    ("IoT / GPS lock", "Smart lock & tracking where supported."),
]
col_w = Inches(3.9); row_h = Inches(1.9)
for i, (h, b) in enumerate(specs):
    col, row = i % 3, i // 3
    x = Inches(0.8) + col * (col_w + Inches(0.2))
    y = Inches(2.2) + row * (row_h + Inches(0.15))
    card = add_round(s, x, y, col_w, row_h, WHITE, radius=0.05)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_round(s, x + Inches(0.3), y + Inches(0.3), Inches(2.2), Inches(0.6), LIGHT, radius=0.3)
    add_text(s, x + Inches(0.3), y + Inches(0.3), Inches(2.2), Inches(0.6),
             [[(h, 17, True, ACCENT2)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.35), y + Inches(1.1), col_w - Inches(0.6), Inches(0.7),
             [[(b, 13, False, SLATE)]], line_spacing=1.05)
footer(s, 9)

# ---------------------------------------------------------------------------
# Slide 10 — Launch Stations
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
add_round(s, Inches(0.5), Inches(0.55), Inches(0.18), Inches(0.5), GLIDE_GREEN)
add_text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.4),
         [[("LAUNCH FOOTPRINT", 12, True, GLIDE_GREEN)]])
add_text(s, Inches(0.8), Inches(0.95), Inches(11.5), Inches(0.7),
         [[("Three stations to start, in Lamphun", 30, True, WHITE)]])
stations = [
    ("Lamphun Tourism Center", "Visitor-facing hub for tourists and locals."),
    ("Lamphun Railway Station", "First/last-mile link to regional rail."),
    ("Storage & Operations Center", "Charging, maintenance, and fleet base."),
]
for i, (h, b) in enumerate(stations):
    x = Inches(0.8) + i * Inches(4.0)
    card = add_round(s, x, Inches(2.4), Inches(3.7), Inches(3.0), RGBColor(0x22, 0x40, 0x3C), radius=0.06)
    add_ellipse(s, x + Inches(1.45), Inches(2.8), Inches(0.8), Inches(0.8), GLIDE_GREEN)
    add_text(s, x + Inches(1.45), Inches(2.82), Inches(0.8), Inches(0.74),
             [[(str(i + 1), 26, True, DARK)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.25), Inches(3.85), Inches(3.2), Inches(0.8),
             [[(h, 17, True, WHITE)]], align=PP_ALIGN.CENTER, line_spacing=1.0)
    add_text(s, x + Inches(0.3), Inches(4.7), Inches(3.1), Inches(0.7),
             [[(b, 13, False, RGBColor(0xC8, 0xE6, 0xDA))]], align=PP_ALIGN.CENTER, line_spacing=1.05)

# ---------------------------------------------------------------------------
# Slide 11 — MVP vs Phase 2
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "Roadmap", "Simple MVP, expandable system")
# two columns
add_round(s, Inches(0.8), Inches(2.2), Inches(5.7), Inches(4.2), WHITE, radius=0.05)
add_rect(s, Inches(0.8), Inches(2.2), Inches(5.7), Inches(0.65), GLIDE_GREEN)
add_text(s, Inches(0.8), Inches(2.22), Inches(5.7), Inches(0.6),
         [[("MVP — ready for launch", 17, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
mvp = [
    "Offline-first staff tablet at stations",
    "Registration with PDPA consent",
    "Staff-assisted rental & return records",
    "QR / PromptPay / manual slip reconciliation",
    "Incident records with photo evidence",
    "Dashboard, exports, SOP support",
]
for i, t in enumerate(mvp):
    y = Inches(3.05) + i * Inches(0.55)
    add_text(s, Inches(1.1), y, Inches(5.2), Inches(0.5),
             [[("•  ", 14, True, ACCENT2), (t, 13, False, SLATE)]])

add_round(s, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.2), WHITE, radius=0.05)
add_rect(s, Inches(6.8), Inches(2.2), Inches(5.7), Inches(0.65), SLATE)
add_text(s, Inches(6.8), Inches(2.22), Inches(5.7), Inches(0.6),
         [[("Phase 2+ — once stable", 17, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
ph2 = [
    "Live GPS trip map & geofencing",
    "Smart-lock integration beyond mobile assist",
    "Real-time battery & charger telemetry",
    "Green Passport, tourism & merchant rewards",
    "Predictive maintenance & dynamic pricing",
    "Carbon-credit-grade ESG reporting",
]
for i, t in enumerate(ph2):
    y = Inches(3.05) + i * Inches(0.55)
    add_text(s, Inches(7.1), y, Inches(5.2), Inches(0.5),
             [[("•  ", 14, True, SLATE), (t, 13, False, SLATE)]])
footer(s, 11)

# ---------------------------------------------------------------------------
# Slide 12 — Technology (vendor-neutral)
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "Technology", "Self-hosted, sovereign by design")
pillars = [
    ("Local-first deployment", "Containerized services run on private cloud or on-premise hardware you control."),
    ("Relational data core", "A robust, standards-based database holds the full domain model with enforced integrity."),
    ("Native mobile + web", "Cross-platform rider app, offline station tablet, and operations console."),
    ("API-driven & open", "Clean service contracts let the city extend and integrate without lock-in."),
]
col_w = Inches(5.9)
for i, (h, b) in enumerate(pillars):
    col, row = i % 2, i // 2
    x = Inches(0.8) + col * (col_w + Inches(0.3))
    y = Inches(2.2) + row * Inches(1.95)
    card = add_round(s, x, y, col_w, Inches(1.7), WHITE, radius=0.05)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_rect(s, x, y, Inches(0.12), Inches(1.7), ACCENT2)
    add_text(s, x + Inches(0.35), y + Inches(0.25), col_w - Inches(0.6), Inches(0.5),
             [[(h, 17, True, DARK)]])
    add_text(s, x + Inches(0.35), y + Inches(0.8), col_w - Inches(0.6), Inches(0.8),
             [[(b, 13, False, SLATE)]], line_spacing=1.05)
footer(s, 12)

# ---------------------------------------------------------------------------
# Slide 13 — Evidence & Compliance
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "Trust & Compliance", "Evidence at every handoff")
left = [
    ("Full audit trail", "Rental, return, payment, damage & override log IDs, actor, time, station."),
    ("Access control", "Role-based permissions; government viewer is strictly read-only."),
]
right = [
    ("Data residency", "Citizen & trip data stays on locally controlled infrastructure."),
    ("PDPA-ready", "Consent capture and identity handling built into registration."),
]
def arch_col(x, items):
    for i, (h, b) in enumerate(items):
        y = Inches(2.2) + i * Inches(1.65)
        card = add_round(s, x, y, Inches(5.7), Inches(1.45), WHITE, radius=0.05)
        card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
        add_ellipse(s, x + Inches(0.25), y + Inches(0.42), Inches(0.6), Inches(0.6), GLIDE_GREEN)
        add_text(s, x + Inches(0.25), y + Inches(0.44), Inches(0.6), Inches(0.56),
                 [[("◈", 18, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(1.05), y + Inches(0.22), Inches(4.4), Inches(0.5),
                 [[(h, 16, True, DARK)]])
        add_text(s, x + Inches(1.05), y + Inches(0.7), Inches(4.4), Inches(0.7),
                 [[(b, 12, False, SLATE)]], line_spacing=1.05)
arch_col(Inches(0.8), left)
arch_col(Inches(6.8), right)
add_text(s, Inches(0.8), Inches(6.2), Inches(11.8), Inches(0.5),
         [[("Compliance and transparency are part of the data model — not bolted on later.", 13, True, ACCENT2)]],
         align=PP_ALIGN.CENTER)
footer(s, 13)

# ---------------------------------------------------------------------------
# Slide 14 — Success Measures
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
add_round(s, Inches(0.5), Inches(0.55), Inches(0.18), Inches(0.5), GLIDE_GREEN)
add_text(s, Inches(0.8), Inches(0.5), Inches(11), Inches(0.4),
         [[("HOW WE MEASURE SUCCESS", 12, True, GLIDE_GREEN)]])
add_text(s, Inches(0.8), Inches(0.95), Inches(11.5), Inches(0.7),
         [[("Targets we hold the MVP to", 30, True, WHITE)]])
measures = [
    ("90%", "fleet available or explainable by maintenance/charging"),
    ("Backup", "mode can start, return & reconcile without stopping service"),
    ("Payment", "integrity: every payment links to a Rental ID, daily reconciliation"),
    ("Safety", "lock: pending-inspection & maintenance bikes cannot be rented"),
    ("Data", "quality: required rental/return/payment/incident fields complete"),
    ("Transparency", "reports expose utilization, revenue, incidents, carbon, audit"),
]
for i, (big, small) in enumerate(measures):
    col, row = i % 3, i // 3
    x = Inches(0.8) + col * (Inches(3.9) + Inches(0.2))
    y = Inches(2.2) + row * Inches(2.0)
    card = add_round(s, x, y, Inches(3.9), Inches(1.8), RGBColor(0x22, 0x40, 0x3C), radius=0.06)
    add_text(s, x + Inches(0.3), y + Inches(0.2), Inches(3.3), Inches(0.6),
             [[(big, 24, True, GLIDE_GREEN)]])
    add_text(s, x + Inches(0.3), y + Inches(0.85), Inches(3.3), Inches(0.85),
             [[(small, 13, False, RGBColor(0xE6, 0xF2, 0xED))]], line_spacing=1.05)

# ---------------------------------------------------------------------------
# Slide 15 — The Ask / Closing
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s)
section_header(s, "The Ask", "Partner with us to mobilize Lamphun")
ask = [
    ("Pilot partners", "Lamphun stations & agencies ready to launch a local e-bike program."),
    ("Investment", "Funding to reach the 31 Oct 2026 MVP and first stable operations."),
    ("Talent", "Engineers and policy leads who care about sovereign public mobility."),
]
for i, (h, b) in enumerate(ask):
    x = Inches(0.8) + i * Inches(4.0)
    card = add_round(s, x, y if False else Inches(2.2), Inches(3.7), Inches(2.8), WHITE, radius=0.06)
    card.line.color.rgb = RGBColor(0xE0, 0xE8, 0xE4); card.line.width = Pt(1)
    add_round(s, x + Inches(1.35), Inches(2.55), Inches(1.0), Inches(1.0), GLIDE_GREEN, radius=0.5)
    add_text(s, x + Inches(1.35), Inches(2.57), Inches(1.0), Inches(0.96),
             [[(str(i + 1), 30, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.3), Inches(3.8), Inches(3.1), Inches(0.5),
             [[(h, 18, True, DARK)]], align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.35), Inches(4.35), Inches(3.0), Inches(0.9),
             [[(b, 13, False, SLATE)]], align=PP_ALIGN.CENTER, line_spacing=1.05)
add_text(s, Inches(0.8), Inches(5.5), Inches(11.8), Inches(1.0),
         [[("Let's build mobility that stays in the community it serves.", 22, True, ACCENT2)]],
         align=PP_ALIGN.CENTER)
footer(s, 15)

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
out = "/home/eiat/projects/e-bicycle/Lamphun-Smart-Green-Mobility.pptx"
prs.save(out)
print("Saved:", out)
print("Slides:", len(prs.slides._sldIdLst))
